"""Create a minimal, clearly synthetic HIRA workbook pair for each comparison badge."""
from pathlib import Path
from datetime import datetime
import argparse,hashlib,json,zipfile
import openpyxl
from openpyxl.styles import Font,PatternFill

ROOT=Path(__file__).resolve().parents[1]
TARGET=ROOT/'docs/testing/hira-badges'
HEADERS=['번호','조제일자','처방기관','조제기관','제품명','약효분류','성분명','약품코드','단위','1회 투약량','1일 투여횟수','총 투약일수']
CASES=[
 ('01_same','기록 동일',1,1,'1','MAINTAINED','SAME_INGREDIENT_SAME_STRENGTH'),
 ('02_changed','복용정보 변경',1,1,'2','DOSE_CHANGED','STRENGTH_OR_DOSAGE_CHANGED'),
 ('03_new','새로 확인',0,1,'1','NEW','NEW_RECORD'),
 ('04_previous_only','이전에만 있음',1,0,'1','NEEDS_CONFIRMATION','NOT_FOUND_IN_LATEST_SOURCE'),
 ('05_hold','추가 확인 필요',1,1,None,None,None),
]


def create(directory=TARGET):
 directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
 manifest=[]
 for key,label,previous,current,dose,status,reason in CASES:
  paths={}
  for side,count in [('previous',previous),('current',current)]:
   path=directory/key/(side+'.xlsx')
   if path.exists():raise FileExistsError('Existing demo files preserved; generate into a new directory.')
   path.parent.mkdir(parents=True,exist_ok=True)
   wb=openpyxl.Workbook();sheet=wb.active;sheet.title='가상 HIRA 시연'
   wb.properties.creator='YAKEUM SYNTHETIC DEMO';wb.properties.title='가상 자료 전용 - '+label
   wb.properties.description='No real personal or patient data. LOCAL_MOCK matching only.'
   wb.properties.created=datetime(2026,10,29)
   sheet.append([f'약이음 가상 시연 자료 · {label} · {"이전" if side=="previous" else "현재"} · 실제 환자 자료 아님'+(' · 등록 기록 없음(헤더만)' if not count else '')])
   sheet.append(HEADERS)
   if count:
    sheet.append([1,'2026-09-01' if side=='previous' else '2026-10-01','가상처방기관','가상조제기관','가상제품','가상분류',None,'000123456','정','1' if side=='previous' else dose,'1','7'])
    sheet.cell(3,8).number_format='@'
   for cell in sheet[2]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24645B')
   sheet.cell(1,1).font=Font(bold=True,color='9A3412');sheet.freeze_panes='A3'
   for index in range(1,13):sheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width=18
   sheet.row_dimensions[1].height=26
   wb.save(path);wb.close();paths[side]=path.relative_to(directory).as_posix()
  manifest.append({'case':key,'label':label,'files':paths,'rows':{'previous':previous,'current':current},'executed':status is not None,'statuses':[] if status is None else [status],'reason_codes':[] if reason is None else [reason],'hold_kind':None if status else 'UNKNOWN_COMPARISON_FIELDS'})
 (directory/'cases.json').write_text(json.dumps({'synthetic_only':True,'matching':'LOCAL_MOCK','baseline_commit':'4423a96','cases':manifest},ensure_ascii=False,indent=2)+'\n')
 lines=[hashlib.sha256(p.read_bytes()).hexdigest()+'  '+p.relative_to(directory).as_posix() for p in sorted(directory.glob('*/*.xlsx'))]
 (directory/'SHA256SUMS').write_text('\n'.join(lines)+'\n')


def package(directory,output):
 directory=Path(directory);output=Path(output)
 if output.resolve().is_relative_to(directory.resolve()):raise ValueError('ZIP must be outside the input kit directory.')
 output.parent.mkdir(parents=True,exist_ok=True)
 with zipfile.ZipFile(output,'w',zipfile.ZIP_DEFLATED) as z:
  z.comment=b'SYNTHETIC DATA ONLY - NO REAL PERSONAL OR PATIENT DATA'
  for p in sorted(directory.rglob('*')):
   if p.is_file():z.write(p,'YAKEUM_HIRA_5_BADGES/'+p.relative_to(directory).as_posix())
 print('Synthetic demo ZIP prepared:',output)


if __name__=='__main__':
 parser=argparse.ArgumentParser();parser.add_argument('--directory',type=Path,default=TARGET);parser.add_argument('--package-only',action='store_true');parser.add_argument('--zip',type=Path)
 args=parser.parse_args()
 if not args.package_only:create(args.directory)
 if args.zip:package(args.directory,args.zip)

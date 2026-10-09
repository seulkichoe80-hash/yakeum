"""B 내부 결과 → 공통 규격(DATA_SCHEMA / API_CONTRACT) 변환 테스트. 외부 API 호출 없음, 가상 데이터만 사용."""
import json
import uuid
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from backend.services.hira.parser import _num, parse_hira_excel
from backend.services.hira.tests.make_fixtures import COMMON_HEADERS, make_common_sample_xlsx
from backend.services.mfds.mock_client import MockMfdsClient
from backend.services.mfds.models import OfficialIngredient, OfficialItem
from backend.services.normalization import cli
from backend.services.normalization.adapter import (
    RECORD_KEYS,
    to_common_result,
    to_common_results,
    to_error_detail,
    validate_common_result,
)
from backend.services.normalization.models import decimal_to_str
from backend.services.normalization.service import normalize_hira_file

FAMILY = "00000000-0000-4000-8000-000000000001"  # 가상 ID (API_CONTRACT Mock 과 동일)
EXPECTED = Path(__file__).parent / "expected_common_sample.json"


def _common(tmp_path, client=None):
    x = make_common_sample_xlsx(tmp_path / "sample_common.xlsx")
    return to_common_results(normalize_hira_file(x, client=client), FAMILY)


# ---- 필드 매핑 --------------------------------------------------------------------------
def test_matched_record_has_all_schema_fields_in_order(tmp_path):
    r = _common(tmp_path)[0]
    assert r["status"] == "MATCHED"
    rec = r["record"]
    assert list(rec) == RECORD_KEYS
    assert rec["family_member_id"] == FAMILY
    assert rec["source_type"] == "HIRA"
    assert rec["record_date"] == rec["dispensing_date"] == "2024-03-05"
    assert rec["prescribing_institution"] == "가나다의원" and rec["dispensing_institution"] == "라마바약국"
    assert rec["product_name_raw"] == "타이레놀정500밀리그램"
    assert rec["product_name_normalized"] == "타이레놀정500밀리그램"  # 공식 제품명(단일 매칭)
    assert rec["drug_code"] == "000123456" and rec["drug_category"] == "해열진통소염제" and rec["unit"] == "정"
    assert (rec["dose_per_administration"], rec["administrations_per_day"], rec["duration_days"]) == ("1", "3", "5")
    assert rec["mfds_product_id"] == "OFFICIAL-0001"  # 식약처 ID. drug_code 와 별개
    assert rec["mfds_product_id"] != rec["drug_code"]
    assert rec["dosage_text_raw"] is None and rec["confidence"] is None and rec["user_confirmed"] is False
    assert r["match"] == {"source": "MFDS_DRUG", "official_id": "OFFICIAL-0001", "match_method": "CODE_EXACT"}


def test_ingredient_fields_and_null_strength_not_guessed(tmp_path):
    ing = _common(tmp_path)[0]["record"]["ingredients"][0]
    assert ing == {
        "ingredient_name_raw": "아세트아미노펜",
        "ingredient_name_normalized": "아세트아미노펜",
        "product_strength": None,  # HIRA 에 함량 컬럼이 없으면 추측하지 않는다
        "product_strength_unit": None,
        "official_ingredient_id": None,
    }


def test_official_ingredient_is_not_presented_as_raw(tmp_path):
    item = OfficialItem("MFDS_DRUG", "OFF-9", "가상정", [OfficialIngredient("가상성분", Decimal("0.50"), "mg", "M0001")])
    p = tmp_path / "a.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["조제일자", "제품명"])
    ws.append(["2024-01-02", "가상정"])
    wb.save(p)
    r = to_common_results(normalize_hira_file(p, client=MockMfdsClient(items=[item])), FAMILY)[0]
    ing = r["record"]["ingredients"][0]
    assert ing["ingredient_name_raw"] is None and ing["ingredient_name_normalized"] == "가상성분"
    assert ing["product_strength"] == "0.5" and ing["product_strength_unit"] == "mg"
    assert ing["official_ingredient_id"] == "M0001"
    assert r["record"]["mfds_product_id"] == "OFF-9" and r["record"]["drug_code"] is None
    assert validate_common_result(r) == []


def test_statuses_and_failed_record_is_null(tmp_path):
    res = _common(tmp_path)
    assert [r["status"] for r in res] == ["MATCHED", "PARTIAL", "PARTIAL", "UNMATCHED", "FAILED", "FAILED"]
    for r in res:
        assert (r["record"] is None) == (r["status"] == "FAILED")
    failed = res[-1]
    assert failed["match"] == {"source": None, "official_id": None, "match_method": None}
    assert failed["errors"][0]["code"] == "REQUIRED_FIELD_MISSING" and failed["errors"][0]["field"] == "product_name_raw"


def test_bad_dispensing_date_fails_without_guessing_record_date(tmp_path):
    r = _common(tmp_path)[4]
    assert r["status"] == "FAILED" and r["record"] is None
    assert any(e["field"] == "dispensing_date" and e["code"] == "REQUIRED_FIELD_MISSING" for e in r["errors"])
    assert any("이상 날짜" in w["message"] for w in r["warnings"])  # 원인 경고는 보존


def test_unmatched_keeps_raw_record_and_null_match(tmp_path):
    r = _common(tmp_path)[3]
    assert r["status"] == "UNMATCHED"
    assert r["match"] == {"source": None, "official_id": None, "match_method": None}
    assert r["record"]["product_name_raw"] == "없는약정" and r["record"]["mfds_product_id"] is None
    assert r["record"]["dose_per_administration"] == "0"  # 0 과 null 구분
    assert any(w["code"] == "OFFICIAL_MATCH_NOT_FOUND" for w in r["warnings"])


def test_multiple_candidates_not_auto_confirmed(tmp_path):
    r = _common(tmp_path)[1]
    assert r["status"] == "PARTIAL" and r["record"]["mfds_product_id"] is None
    assert r["match"]["official_id"] is None and r["match"]["match_method"] == "NAME_EXACT"
    assert r["record"]["dose_per_administration"] == "0.5"


def test_legacy_single_institution_column_is_not_guessed(tmp_path):
    p = tmp_path / "legacy.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["조제일자", "요양기관명", "약품명"])
    ws.append(["2024-01-02", "가나다의원", "가상정"])
    wb.save(p)
    rec = to_common_results(normalize_hira_file(p), FAMILY)[0]["record"]
    assert rec["prescribing_institution"] is None and rec["dispensing_institution"] is None


# ---- 식별자 ----------------------------------------------------------------------------
def test_ids_are_uuid_deterministic_and_duplicates_are_distinct(tmp_path):
    a = _common(tmp_path)
    assert a == _common(tmp_path)  # 같은 입력 → 같은 출력
    ids = [r["record"]["record_id"] for r in a if r["record"]]
    assert len(set(ids)) == len(ids)
    for i in ids:
        uuid.UUID(i)
    src = [r["record"]["source_record_id"] for r in a if r["record"]]
    assert src[1] != src[2] and src[2] == src[1] + "-2"  # 동일 내용 행 구분
    assert all(s.startswith("hira:") for s in src)


def test_source_record_id_independent_of_row_position(tmp_path):
    a = _common(tmp_path)[0]["record"]["source_record_id"]
    p = tmp_path / "shift.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    for _ in range(5):
        ws.append(["안내"])
    ws.append(COMMON_HEADERS)
    ws.append([1, "2024-03-05", "가나다의원", "라마바약국", "타이레놀정500밀리그램", "해열진통소염제", "아세트아미노펜", "000123456", "정", 1, 3, 5])
    wb.save(p)
    b = to_common_results(normalize_hira_file(p), FAMILY)[0]["record"]["source_record_id"]
    assert a == b


def test_family_member_id_is_required_and_never_guessed(tmp_path):
    res = normalize_hira_file(make_common_sample_xlsx(tmp_path / "s.xlsx"))
    for bad in (None, "", "not-a-uuid"):
        with pytest.raises(ValueError):
            to_common_result(res[0], bad)


def test_record_id_factory_override(tmp_path):
    res = normalize_hira_file(make_common_sample_xlsx(tmp_path / "s.xlsx"))
    fixed = "10000000-0000-4000-8000-000000000009"
    out = to_common_result(res[0], FAMILY, record_id_factory=lambda f, t, s: fixed)
    assert out["record"]["record_id"] == fixed


# ---- Decimal ---------------------------------------------------------------------------
def test_no_float_in_parser_or_output(tmp_path):
    for row in parse_hira_excel(make_common_sample_xlsx(tmp_path / "s.xlsx")):
        d = row.record.dosage
        for v in (d.dose_per_time, d.times_per_day, d.total_days):
            assert v is None or isinstance(v, Decimal)

    def walk(o):
        assert not isinstance(o, float)
        if isinstance(o, dict):
            [walk(v) for v in o.values()]
        if isinstance(o, list):
            [walk(v) for v in o]

    walk(_common(tmp_path))


def test_decimal_precision_preserved():
    assert _num(0.1)[0] == Decimal("0.1")  # float 2진 오차 없이
    assert _num("0.30")[0] == Decimal("0.30")
    assert decimal_to_str(Decimal("0.30")) == "0.3"
    assert decimal_to_str(Decimal("1E+2")) == "100" and decimal_to_str(Decimal("2.0")) == "2"
    assert decimal_to_str(Decimal("0")) == "0"
    v, w = _num("1/3")
    assert v is not None and w and "근사" in w  # 근사 환산은 경고
    assert _num("1/2") == (Decimal("0.5"), None)
    assert _num("abc")[0] is None


# ---- 매칭 코드 체계 -------------------------------------------------------------------------
def test_hira_drug_code_is_not_compared_with_official_product_id():
    # official_id 가 HIRA 약품코드와 우연히 같아도 다른 체계이므로 CODE_EXACT 로 보지 않는다
    c = MockMfdsClient(items=[OfficialItem("MFDS_DRUG", "000123456", "다른약", [])])
    assert c.find_by_code("000123456") == []


# ---- ErrorDetail -----------------------------------------------------------------------
@pytest.mark.parametrize("msg,code,field", [
    ("헤더 행을 찾을 수 없음 (탐색 시트: ['a'])", "INVALID_FILE_FORMAT", None),
    ("약품명이 비어 있어 행을 해석할 수 없음", "REQUIRED_FIELD_MISSING", "product_name_raw"),
    ("공식 데이터에서 찾지 못함 (원본 값 보존)", "OFFICIAL_MATCH_NOT_FOUND", None),
    ("후보 2건 (NAME_EXACT): A, B", "OFFICIAL_MATCH_NOT_FOUND", None),
    ("이상 날짜: '2024-13-45'", "NORMALIZATION_FAILED", "record_date"),
    ("dose_per_time: 숫자로 해석 불가: 'x'", "NORMALIZATION_FAILED", "dose_per_administration"),
    ("약품코드가 숫자 셀이라 앞자리 0이 손실되었을 수 있음", "NORMALIZATION_FAILED", "drug_code"),
    ("성분 정보 없음 (공식 데이터 매칭으로 보완 필요)", "NORMALIZATION_FAILED", "ingredients"),
    ("중복 행: 6행과 동일한 내용", "NORMALIZATION_FAILED", None),
])
def test_error_detail_mapping(msg, code, field):
    assert to_error_detail(msg) == {"code": code, "message": msg, "field": field}


# ---- 계약 점검 / 샘플 출력 / CLI --------------------------------------------------------
def test_all_outputs_pass_contract_validation_and_are_json(tmp_path):
    for r in _common(tmp_path):
        assert validate_common_result(r) == [], r
        json.dumps(r, ensure_ascii=False)


def test_validator_catches_violations(tmp_path):
    r = _common(tmp_path)[0]
    r["record"]["dose_per_administration"] = 1.5
    r["record"]["drug_code"] = 123
    r["errors"] = ["문자열 오류"]
    probs = validate_common_result(r)
    assert any("Decimal" in p for p in probs)
    assert any("drug_code" in p for p in probs)
    assert any("ErrorDetail" in p for p in probs)


def test_fields_needed_by_reconciliation_are_present(tmp_path):
    rec = _common(tmp_path)[0]["record"]
    for k in ("family_member_id", "record_id", "mfds_product_id", "dose_per_administration"):
        assert rec[k] not in (None, "")
    assert rec["ingredients"]


def test_sample_output_matches_expected_json(tmp_path):
    assert json.loads(json.dumps(_common(tmp_path))) == json.loads(EXPECTED.read_text(encoding="utf-8"))


def test_cli_common_output(tmp_path):
    x = make_common_sample_xlsx(tmp_path / "s.xlsx")
    out = tmp_path / "out.json"
    assert cli.main([str(x), "--common", "--family-member-id", FAMILY, "--out", str(out)]) == 0
    data = json.loads(out.read_text(encoding="utf-8"))
    assert len(data) == 6 and data[0]["record"]["family_member_id"] == FAMILY
    with pytest.raises(SystemExit):
        cli.main([str(x), "--common"])  # family_member_id 없이는 거부

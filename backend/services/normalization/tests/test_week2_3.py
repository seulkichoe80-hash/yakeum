import json
from datetime import datetime

import openpyxl

from backend.services.hira.parser import parse_date, parse_hira_excel
from backend.services.hira.tests.make_fixtures import make_sample_xlsx
from backend.services.mfds.mock_client import MockMfdsClient
from backend.services.mfds.models import OfficialItem
from backend.services.mfds.real_client import RealMfdsClient
from backend.services.normalization import cli
from backend.services.normalization.service import normalize_hira_file, normalize_rows, summarize


class BoomClient(MockMfdsClient):
    def find_by_code(self, code):
        raise RuntimeError("boom")

    def search_by_name(self, name, limit=50):
        raise RuntimeError("boom")


def _xlsx(tmp_path, rows, name="t.xlsx", header=("진료일자", "약품명", "약품코드", "1회 투약량")):
    p = tmp_path / name
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(header))
    for r in rows:
        ws.append(r)
    wb.save(p)
    return p


def test_lookup_failure_is_unmatched_with_failure_warning(tmp_path):
    c = RealMfdsClient(api_key="K", retries=0, transport=lambda u, p, t: (_ for _ in ()).throw(TimeoutError("t")), sleep=lambda s: None)
    r = normalize_hira_file(make_sample_xlsx(tmp_path / "s.xlsx"), client=c)
    assert [x.status for x in r] == ["UNMATCHED", "UNMATCHED", "UNMATCHED", "UNMATCHED", "FAILED"]
    assert any("조회 실패로 매칭하지 못함" in w for w in r[0].warnings)
    assert not any("찾지 못함" in w for w in r[0].warnings)  # "없음"과 "실패"를 구분
    assert r[0].record.product_name.raw == "타이레놀정500밀리그램"  # 원본 보존


def test_client_exception_does_not_stop_other_rows(tmp_path):
    rows = parse_hira_excel(make_sample_xlsx(tmp_path / "s.xlsx"))
    res = normalize_rows(rows, client=BoomClient())
    assert len(res) == 5
    assert all(r.status in ("UNMATCHED", "FAILED") for r in res)
    assert any("조회 실패" in w for w in res[0].warnings)


def test_single_fuzzy_candidate_is_matched():
    client = MockMfdsClient(items=[OfficialItem("MFDS_DRUG", "1", "뮤코펙트정", [])])
    from backend.services.normalization.matcher import match
    out = match(client, None, "뮤코펙토정", 0.7)
    assert out.method == "NAME_FUZZY" and len(out.candidates) == 1


def test_matched_without_any_ingredients_is_partial(tmp_path):
    client = MockMfdsClient(items=[OfficialItem("MFDS_HEALTH_FOOD", "9", "비타민", [])])
    p = _xlsx(tmp_path, [["2024-01-01", "비타민", None, 1]])
    r = normalize_hira_file(p, client=client)[0]
    assert r.status == "PARTIAL" and r.match.official_id == "9"


def test_invalid_threshold_rejected(tmp_path):
    import pytest
    with pytest.raises(ValueError):
        normalize_rows([], config={"fuzzy_threshold": 1.5})


def test_summarize(tmp_path):
    r = normalize_hira_file(make_sample_xlsx(tmp_path / "s.xlsx"))
    assert summarize(r) == {"total": 5, "MATCHED": 2, "PARTIAL": 1, "UNMATCHED": 1, "FAILED": 1}


def test_date_string_with_time_and_fraction_dose(tmp_path):
    assert parse_date("2024-03-05 00:00:00")[0] == "2024-03-05"
    assert parse_date("2024-03-05T09:30:00")[0] == "2024-03-05"
    p = _xlsx(tmp_path, [["2024-03-05 00:00:00", "가나정", "001", "1/2"]])
    rec = parse_hira_excel(p)[0].record
    assert rec.date == "2024-03-05" and rec.dosage.dose_per_time == 0.5


def test_duplicate_rows_warn(tmp_path):
    p = _xlsx(tmp_path, [["2024-01-01", "가나정", "001", 1], ["2024-01-01", "가나정", "001", 1]])
    rows = parse_hira_excel(p)
    assert not any("중복" in w for w in rows[0].warnings)
    assert any("중복 행" in w for w in rows[1].warnings)


def test_header_found_on_non_active_sheet(tmp_path):
    p = tmp_path / "multi.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "표지"
    wb.active.append(["안내문"])
    ws = wb.create_sheet("내역")
    ws.append(["진료일자", "약품명", "약품코드"])
    ws.append(["2024-01-01", "가나정", "001"])
    wb.save(p)
    rows = parse_hira_excel(p)
    assert len(rows) == 1 and rows[0].record.source_ref["sheet"] == "내역" and not rows[0].errors


def test_header_with_nbsp_and_spaces(tmp_path):
    p = _xlsx(tmp_path, [["2024-01-01", "가나정", "001", 1]], header=("진료 일자", " 약품 명 ", "보험코드", "1회 투약량"))
    assert not parse_hira_excel(p)[0].errors


def test_cli_writes_json_and_summary(tmp_path, capsys):
    x = make_sample_xlsx(tmp_path / "s.xlsx")
    out = tmp_path / "out.json"
    assert cli.main([str(x), "--out", str(out)]) == 0
    data = json.loads(out.read_text(encoding="utf-8"))
    assert len(data) == 5 and data[0]["match"]["match_method"] == "CODE_EXACT"
    assert '"total": 5' in capsys.readouterr().err


def test_deterministic_output(tmp_path):
    x = make_sample_xlsx(tmp_path / "s.xlsx")
    a = [r.to_dict() for r in normalize_hira_file(x)]
    b = [r.to_dict() for r in normalize_hira_file(x)]
    assert a == b

from datetime import datetime

import openpyxl

from backend.services.hira.parser import parse_date, parse_hira_excel
from backend.services.hira.tests.make_fixtures import make_sample_xlsx


def _rows(tmp_path):
    return parse_hira_excel(make_sample_xlsx(tmp_path / "sample_hira.xlsx"))


def test_header_detected_below_title_and_blank_row_skipped(tmp_path):
    rows = _rows(tmp_path)
    assert len(rows) == 5  # 빈 행 제외
    assert rows[0].record.source_ref["row"] == 6


def test_normal_row(tmp_path):
    r = _rows(tmp_path)[0]
    assert r.errors == []
    rec = r.record
    assert rec.date == "2024-03-05"
    assert rec.drug_code == "000123456"
    assert rec.product_name.raw == "타이레놀정500밀리그램"
    assert rec.ingredients[0].amount == 500 and rec.ingredients[0].unit == "mg"
    assert (rec.dosage.dose_per_time, rec.dosage.times_per_day, rec.dosage.total_days) == (1, 3, 5)


def test_leading_zero_code_stays_string(tmp_path):
    assert isinstance(_rows(tmp_path)[0].record.drug_code, str)
    assert _rows(tmp_path)[0].record.drug_code.startswith("000")


def test_zero_dose_differs_from_missing(tmp_path):
    rows = _rows(tmp_path)
    assert rows[1].record.dosage.dose_per_time == 0
    assert rows[3].record.dosage.dose_per_time is None


def test_bad_date_is_null_with_warning(tmp_path):
    r = _rows(tmp_path)[2]
    assert r.record.date is None and r.record.date_raw == "2024-13-45"
    assert any("이상 날짜" in w for w in r.warnings)


def test_missing_name_fails_only_that_row(tmp_path):
    rows = _rows(tmp_path)
    assert rows[-1].errors and not rows[0].errors


def test_numeric_code_cell_warns(tmp_path):
    p = tmp_path / "n.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["진료일자", "약품명", "약품코드"])
    ws.append([datetime(2024, 1, 2), "테스트정", 123456])
    wb.save(p)
    r = parse_hira_excel(p)[0]
    assert r.record.drug_code == "123456"
    assert any("앞자리 0" in w for w in r.warnings)


def test_alias_headers_and_merged_date(tmp_path):
    p = tmp_path / "m.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["조제일자", "처방약품명", "보험코드"])
    ws.append(["20240110", "가나정", "001"])
    ws.append([None, "다라정", "002"])
    ws.merge_cells("A2:A3")
    wb.save(p)
    rows = parse_hira_excel(p)
    assert [r.record.date for r in rows] == ["2024-01-10", "2024-01-10"]
    assert any("병합" in w for w in rows[1].warnings)


def test_no_header_returns_failed_row(tmp_path):
    p = tmp_path / "x.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["아무", "내용"])
    wb.save(p)
    rows = parse_hira_excel(p)
    assert len(rows) == 1 and rows[0].errors


def test_unreadable_file(tmp_path):
    p = tmp_path / "bad.xlsx"
    p.write_text("not excel")
    assert parse_hira_excel(p)[0].errors


def test_parse_date_formats():
    assert parse_date("2024.3.5")[0] == "2024-03-05"
    assert parse_date("20240305")[0] == "2024-03-05"
    assert parse_date("0000-00-00")[0] is None

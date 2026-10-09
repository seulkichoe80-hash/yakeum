"""HIRA 진료내역 Excel 규칙 기반 파서 (LLM/AI 미사용, 결정적).

추측하지 않는다: 모르면 None + warnings. 한 행이 실패해도 전체는 계속 진행한다.
"""
import logging
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
import yaml

from backend.services.normalization.models import (
    Dosage,
    Ingredient,
    MedicationRecord,
    ParsedRow,
    TextField,
)
from backend.services.normalization.name_normalizer import normalize_name

log = logging.getLogger(__name__)
DEFAULT_MAPPING = Path(__file__).parent / "column_mapping.yaml"

_STRENGTH_RE = re.compile(r"^\s*(\d[\d,]*(?:\.\d+)?)\s*([A-Za-zμ㎍㎎㎖%가-힣]*)\s*$")
_NUM_PREFIX_RE = re.compile(r"^\s*(\d[\d,]*(?:\.\d+)?)")
_FRACTION_RE = re.compile(r"^\s*(\d+)\s*/\s*(\d+)\s*[A-Za-z가-힣]*\s*$")
_DATE_RES = [
    re.compile(r"^(\d{4})[-./](\d{1,2})[-./](\d{1,2})\.?$"),
    re.compile(r"^(\d{4})(\d{2})(\d{2})$"),
]


def load_mapping(path: Optional[Path] = None) -> Dict[str, Any]:
    return yaml.safe_load(Path(path or DEFAULT_MAPPING).read_text(encoding="utf-8"))


def _key(v: Any) -> str:
    return re.sub(r"\s+", "", str(v)).lower() if v is not None else ""


def _blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def _text(v: Any) -> Optional[str]:
    return None if _blank(v) else str(v).strip()


def _num(v: Any) -> Tuple[Optional[float], Optional[str]]:
    """(값, 경고). 0 은 0 으로, 빈 값은 None 으로 구분."""
    if _blank(v):
        return None, None
    if isinstance(v, bool):
        return None, f"숫자로 해석 불가: {v!r}"
    if isinstance(v, (int, float)):
        n = float(v)
    else:
        fr = _FRACTION_RE.match(str(v))
        if fr and float(fr.group(2)) != 0:
            n = float(fr.group(1)) / float(fr.group(2))
            return (int(n) if n == int(n) else n), None
        m = _NUM_PREFIX_RE.match(str(v))
        if not m:
            return None, f"숫자로 해석 불가: {v!r}"
        n = float(m.group(1).replace(",", ""))
    return (int(n) if n == int(n) else n), None


def _grid(ws) -> Tuple[List[List[Any]], set]:
    """병합 셀을 값으로 채운 2차원 리스트와, 병합으로 채워진 좌표 집합을 반환."""
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    filled = set()
    for rng in ws.merged_cells.ranges:
        v = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    while len(grid) < r:
                        grid.append([])
                    row = grid[r - 1]
                    row.extend([None] * (c - len(row)))
                    row[c - 1] = v
                    filled.add((r - 1, c - 1))
    return grid, filled


def _detect_header(grid, mapping) -> Tuple[Optional[int], Dict[str, int]]:
    aliases = {f: {_key(a) for a in al} for f, al in mapping["fields"].items()}
    best, best_cols, best_n = None, {}, 0
    for i, row in enumerate(grid[: mapping.get("header_scan_rows", 30)]):
        cols: Dict[str, int] = {}
        for j, cell in enumerate(row):
            k = _key(cell)
            for f, al in aliases.items():
                if k and k in al and f not in cols:
                    cols[f] = j
        if len(cols) > best_n:
            best, best_cols, best_n = i, cols, len(cols)
    if best is None or best_n < mapping.get("min_header_matches", 2):
        return None, {}
    return best, best_cols


def parse_date(v: Any, year_range=(1990, 2100)) -> Tuple[Optional[str], Optional[str]]:
    """(YYYY-MM-DD, 경고). 해석 불가/이상 날짜는 None + 경고."""
    if _blank(v):
        return None, None
    d: Optional[date] = None
    if isinstance(v, datetime):
        d = v.date()
    elif isinstance(v, date):
        d = v
    else:
        s = str(int(v)) if isinstance(v, float) and v == int(v) else str(v).strip()
        s = re.split(r"[ T]", s, maxsplit=1)[0]  # "2024-03-05 00:00:00" 같은 시각 부분 제거
        for rx in _DATE_RES:
            m = rx.match(s)
            if m:
                try:
                    d = date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
                except ValueError:
                    d = None
                break
    if d is None or not (year_range[0] <= d.year <= year_range[1]):
        return None, f"이상 날짜: {v!r}"
    return d.isoformat(), None


def _split(s: Optional[str]) -> List[str]:
    return [p.strip() for p in re.split(r"[/,+]", s) if p.strip()] if s else []


def _ingredients(ing_raw, str_raw, warns) -> List[Ingredient]:
    names, strengths = _split(ing_raw), _split(str_raw)
    if not names:
        if str_raw:
            warns.append("함량은 있으나 성분명이 없어 성분 정보를 비워 둠")
        else:
            warns.append("성분 정보 없음 (공식 데이터 매칭으로 보완 필요)")
        return []
    if len(strengths) not in (0, 1, len(names)):
        warns.append(f"성분 수({len(names)})와 함량 수({len(strengths)}) 불일치: 함량 미할당")
        strengths = []
    if len(strengths) == 1 and len(names) > 1:
        warns.append("성분이 여러 개인데 함량이 1개: 함량 미할당")
        strengths = []
    out = []
    for i, n in enumerate(names):
        s = strengths[i] if strengths else None
        amount, unit = None, None
        if s:
            m = _STRENGTH_RE.match(s)
            if m:
                amount = float(m.group(1).replace(",", ""))
                amount = int(amount) if amount == int(amount) else amount
                unit = m.group(2) or None
            else:
                warns.append(f"함량 해석 불가: {s!r}")
        out.append(Ingredient(name=TextField(n, normalize_name(n)), amount=amount, unit=unit, amount_raw=s))
    return out


def _row_to_record(raw: Dict[str, Any], mapping, ref, warns, errs) -> MedicationRecord:
    yr = tuple(mapping.get("year_range", (1990, 2100)))
    rec = MedicationRecord(source_ref=ref, raw_row={k: (None if _blank(v) else str(v)) for k, v in raw.items()})

    name = _text(raw.get("product_name"))
    if name is None:
        errs.append("약품명이 비어 있어 행을 해석할 수 없음")
    rec.product_name = TextField(name, normalize_name(name))

    code = raw.get("drug_code")
    if not _blank(code):
        if isinstance(code, (int, float)) and not isinstance(code, bool):
            rec.drug_code = str(int(code))
            warns.append("약품코드가 숫자 셀이라 앞자리 0이 손실되었을 수 있음")
        else:
            rec.drug_code = str(code).strip()  # 문자열 그대로 (앞자리 0 보존)
    else:
        warns.append("약품코드 없음")

    rec.date_raw = _text(raw.get("date"))
    rec.date, w = parse_date(raw.get("date"), yr)
    if w:
        warns.append(w)
    elif rec.date is None:
        warns.append("날짜 없음")

    rec.institution = _text(raw.get("institution"))
    rec.ingredients = _ingredients(_text(raw.get("ingredient")), _text(raw.get("strength")), warns)

    d = Dosage(raw={k: _text(raw.get(k)) for k in ("dose_per_time", "times_per_day", "total_days")})
    for fld in ("dose_per_time", "times_per_day", "total_days"):
        val, w = _num(raw.get(fld))
        if w:
            warns.append(f"{fld}: {w}")
        if fld == "total_days" and val is not None:
            if val != int(val):
                warns.append(f"total_days: 정수가 아님: {val!r}")
                val = None
            else:
                val = int(val)
        setattr(d, fld, val)
    rec.dosage = d
    return rec


def parse_hira_excel(path, mapping_path: Optional[Path] = None, sheet: Optional[str] = None) -> List[ParsedRow]:
    """Excel 파일 → ParsedRow 리스트. 파일/헤더 단위 실패는 예외 대신 FAILED 행 1개로 반환."""
    mapping = load_mapping(mapping_path)
    path = Path(path)
    base_ref = {"file": path.name, "sheet": None, "row": None}
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:  # noqa: BLE001 - 파일을 못 읽는 경우도 결과로 기록
        return [ParsedRow(MedicationRecord(source_ref=base_ref), errors=[f"파일을 읽을 수 없음: {e}"])]

    # 시트 지정이 없으면 활성 시트부터 보고, 헤더가 없으면 다른 시트도 탐색
    sheets = [wb[sheet]] if sheet else [wb.active] + [w for w in wb.worksheets if w is not wb.active]
    ws = grid = merged = None
    h, cols = None, {}
    for cand in sheets:
        try:
            g, mg = _grid(cand)
        except Exception as e:  # noqa: BLE001
            log.warning("sheet %s unreadable: %s", cand.title, e)
            continue
        h, cols = _detect_header(g, mapping)
        if h is not None:
            ws, grid, merged = cand, g, mg
            break
    if h is None:
        return [ParsedRow(MedicationRecord(source_ref=base_ref), errors=[f"헤더 행을 찾을 수 없음 (탐색 시트: {[w.title for w in sheets]})"])]
    base_ref["sheet"] = ws.title

    missing = [f for f in mapping.get("required_fields", []) if f not in cols]
    if missing:
        return [ParsedRow(MedicationRecord(source_ref=base_ref), errors=[f"필수 컬럼 없음: {missing}"])]

    rows: List[ParsedRow] = []
    seen: Dict[Any, int] = {}
    for i in range(h + 1, len(grid)):
        row = grid[i]
        if all(_blank(c) for c in row):
            continue  # 빈 행
        ref = {**base_ref, "row": i + 1}
        warns: List[str] = []
        errs: List[str] = []
        raw: Dict[str, Any] = {}
        for f, j in cols.items():
            raw[f] = row[j] if j < len(row) else None
            if (i, j) in merged and not _blank(raw[f]):
                warns.append(f"{f}: 병합 셀 값을 채워 사용")
        try:
            rec = _row_to_record(raw, mapping, ref, warns, errs)
        except Exception as e:  # noqa: BLE001 - 한 행 실패가 전체를 막지 않게
            rec = MedicationRecord(source_ref=ref, raw_row={k: _text(v) for k, v in raw.items()})
            errs.append(f"행 처리 중 예외: {e}")
        key = (rec.date_raw, rec.drug_code or rec.product_name.raw, rec.institution, tuple(sorted(rec.dosage.raw.items())))
        if not errs and key in seen:
            warns.append(f"중복 행: {seen[key]}행과 동일한 내용")
        elif not errs:
            seen[key] = i + 1
        rows.append(ParsedRow(rec, warns, errs))
    log.info("parsed %s: %d rows", path.name, len(rows))
    return rows

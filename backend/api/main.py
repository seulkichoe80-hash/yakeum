"""Local-only FastAPI entrypoint for the first B-to-D-to-C Vertical Slice."""

from __future__ import annotations

import tempfile
import zipfile
from pathlib import Path, PurePosixPath
from typing import Annotated
from uuid import UUID

import openpyxl
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.exception_handlers import request_validation_exception_handler
from fastapi.responses import JSONResponse

from backend.api.pipeline import process_hira_comparison
from backend.api.schemas import VerticalSliceResponse, HiraV2Response
from backend.api.photo_demo import router as photo_demo_router
from backend.api.photo_upload import router as photo_upload_router
from backend.api.dev_account import router as dev_account_router
from backend.api.account_hira import router as account_hira_router

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ZIP_ENTRIES = 2_048
MAX_WORKBOOK_ROWS = 10_000
MAX_WORKBOOK_COLUMNS = 128
COPY_CHUNK_BYTES = 1024 * 1024


class UploadRejected(ValueError):
    def __init__(self, status_code: int, detail: str) -> None:
        super().__init__(detail)
        self.status_code = status_code
        self.detail = detail


app = FastAPI(
    title="Yakeum B-D-C Integration API",
    version="0.1.0",
    description=(
        "Mock integration and explicitly enabled development-only accounts/SQLite. "
        "No live medical API or production authentication/storage."
    ),
)
app.include_router(photo_demo_router)
app.include_router(photo_upload_router)
app.include_router(dev_account_router)
app.include_router(account_hira_router)


@app.exception_handler(RequestValidationError)
async def safe_validation_error(request, error):
    if request.url.path.startswith('/api/dev-account'):
        # Do not reflect passwords or uploaded form values in validation errors.
        return JSONResponse(status_code=422, content={"detail": "입력 형식과 필수 항목을 확인해 주세요."})
    return await request_validation_exception_handler(request, error)


@app.middleware("http")
async def private_account_responses(request, call_next):
    response = await call_next(request)
    if request.url.path.startswith("/api/dev-account"):
        response.headers["Cache-Control"] = "no-store"
    return response


async def _reconcile_hira_files(
    family_member_id: Annotated[UUID, Form()],
    previous_file: Annotated[UploadFile, File()],
    current_file: Annotated[UploadFile, File()],
    *, version: int = 1,
) -> VerticalSliceResponse | HiraV2Response:
    """Process explicitly labelled previous/current virtual HIRA XLSX files."""

    try:
        with tempfile.TemporaryDirectory(prefix="yakeum-upload-") as directory:
            root = Path(directory)
            previous_path = root / "previous.xlsx"
            current_path = root / "current.xlsx"
            await _save_and_validate(previous_file, previous_path)
            await _save_and_validate(current_file, current_path)
            result = process_hira_comparison(
                previous_path,
                current_path,
                str(family_member_id),
                version=version,
            )
            return (HiraV2Response if version == 2 else VerticalSliceResponse).model_validate(result)
    except UploadRejected as error:
        raise HTTPException(
            status_code=error.status_code,
            detail=error.detail,
        ) from error
    except ValueError as error:
        raise HTTPException(
            status_code=422,
            detail="입력 자료의 가족 또는 계약 필드를 확인할 수 없습니다.",
        ) from error
    finally:
        await previous_file.close()
        await current_file.close()


@app.post("/api/v1/reconciliations/hira", response_model=VerticalSliceResponse)
async def reconcile_hira_files(
    family_member_id: Annotated[UUID, Form()],
    previous_file: Annotated[UploadFile, File()],
    current_file: Annotated[UploadFile, File()],
):
    return await _reconcile_hira_files(family_member_id, previous_file, current_file)


@app.post("/api/v2/reconciliations/hira", response_model=HiraV2Response)
async def reconcile_hira_v2(
    family_member_id: Annotated[UUID, Form()],
    previous_file: Annotated[UploadFile, File()],
    current_file: Annotated[UploadFile, File()],
):
    return await _reconcile_hira_files(family_member_id, previous_file, current_file, version=2)


async def _save_and_validate(upload: UploadFile, destination: Path) -> None:
    filename = upload.filename or ""
    if Path(filename).suffix.casefold() != ".xlsx":
        raise UploadRejected(415, "XLSX 파일만 업로드할 수 있습니다.")

    size = 0
    with destination.open("xb") as target:
        while chunk := await upload.read(COPY_CHUNK_BYTES):
            size += len(chunk)
            if size > MAX_FILE_BYTES:
                raise UploadRejected(413, "파일당 최대 크기는 10MB입니다.")
            target.write(chunk)

    if size == 0:
        raise UploadRejected(400, "빈 XLSX 파일은 처리할 수 없습니다.")
    _validate_xlsx_structure(destination)


def _validate_xlsx_structure(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            names = {entry.filename for entry in entries}
            required = {"[Content_Types].xml", "xl/workbook.xml"}
            if not required.issubset(names) or not any(
                name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
                for name in names
            ):
                raise UploadRejected(400, "유효한 XLSX 구조를 확인할 수 없습니다.")
            if len(entries) > MAX_ZIP_ENTRIES:
                raise UploadRejected(400, "XLSX 내부 항목 수가 처리 한도를 초과합니다.")

            uncompressed_size = 0
            for entry in entries:
                member = PurePosixPath(entry.filename)
                if member.is_absolute() or ".." in member.parts:
                    raise UploadRejected(400, "안전하지 않은 XLSX 내부 경로입니다.")
                uncompressed_size += entry.file_size
                if uncompressed_size > MAX_UNCOMPRESSED_BYTES:
                    raise UploadRejected(
                        400,
                        "XLSX 압축 해제 크기가 처리 한도를 초과합니다.",
                    )
    except (zipfile.BadZipFile, OSError) as error:
        raise UploadRejected(400, "손상되었거나 유효하지 않은 XLSX 파일입니다.") from error

    try:
        workbook = openpyxl.load_workbook(
            path,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        try:
            total_rows = 0
            for worksheet in workbook.worksheets:
                total_rows += worksheet.max_row
                if (
                    total_rows > MAX_WORKBOOK_ROWS
                    or worksheet.max_column > MAX_WORKBOOK_COLUMNS
                ):
                    raise UploadRejected(
                        400,
                        "XLSX 행 또는 열 수가 처리 한도를 초과합니다.",
                    )
        finally:
            workbook.close()
    except UploadRejected:
        raise
    except Exception as error:  # openpyxl emits several format-specific errors
        raise UploadRejected(400, "손상되었거나 유효하지 않은 XLSX 파일입니다.") from error


# Register after all API routes so static serving cannot shadow them.
from backend.api.web_runtime import configure_web
configure_web(app)

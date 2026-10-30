"""No live API. Synthetic XML -> owned photo/HIRA processing and persistence."""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import xml.etree.ElementTree as ET

import pytest
from fastapi.testclient import TestClient

from backend.api.main import app
from backend.api import mfds_admin
from backend.services.mfds import probe
from backend.services.mfds.live_matching import LiveLookup, verify_result
from backend.tests.test_dev_account import account, family, processed, save, upload, ORIGIN
from backend.tests.test_photo_demo import draft
from backend.tests.test_b_d_c_api import _row,_xlsx


@pytest.fixture(autouse=True)
def environment(monkeypatch,tmp_path):
    for k,v in {'YAKEUM_DEV_ACCOUNT_ENABLED':'1','YAKEUM_APP_MODE':'development','YAKEUM_STORAGE_BACKEND':'sqlite',
                'YAKEUM_DEV_ORIGIN':ORIGIN,'YAKEUM_DEV_DB_PATH':str(tmp_path/'live.sqlite3'),
                'YAKEUM_MFDS_MATCH_ENABLED':'1','MFDS_API_KEY':'SYNTHETIC-SECRET'}.items():monkeypatch.setenv(k,v)
    monkeypatch.setattr(probe,'fetch_xml',lambda *a:pytest.fail('external request forbidden'))


def xml(case='normal'):
    root=ET.Element('response');h=ET.SubElement(root,'header');ET.SubElement(h,'resultCode').text='00'
    body=ET.SubElement(root,'body')
    rows=0 if case=='empty' else 2 if case=='multiple' else 3 if case=='partial' else 1
    for k,v in [('pageNo',1),('numOfRows',3),('totalCount',7 if case=='partial' else rows)]:ET.SubElement(body,k).text=str(v)
    items=ET.SubElement(body,'items')
    for index in range(rows):
        row=ET.SubElement(items,'item')
        values={'ITEM_SEQ':f'SYNTHETIC-{index}','ITEM_NAME':'가상알파정',
                'MAIN_ITEM_INGR':'[SYNTHETIC-ING]가상알파성분','EDI_CODE':'',
                'MATERIAL_NAME':'총량 : 1정200밀리그램|성분명 : 가상알파성분|분량 : 100|단위 : 밀리그램|규격 : 가상규격|성분정보 :'}
        if case=='unknown':values.pop('MATERIAL_NAME')
        if case=='mismatch':values['MATERIAL_NAME']=values['MATERIAL_NAME'].replace('분량 : 100','분량 : 200')
        if case=='granule':values['MATERIAL_NAME']=values['MATERIAL_NAME'].replace('성분정보 :','성분정보 : 가상알파성분으로서 100 밀리그램')
        if case=='name':values['ITEM_NAME']='가상알파정200mg'
        if case=='ingredient':values['MAIN_ITEM_INGR']='[SYNTHETIC]다른성분'
        for k,v in values.items():ET.SubElement(row,k).text=v
    return ET.tostring(root,encoding='unicode')


def fake(monkeypatch,case):
    calls=[]
    def transport(key,name,page,rows):
        calls.append((name,page,rows))
        assert key=='SYNTHETIC-SECRET'
        if case=='error':raise TimeoutError('SYNTHETIC-SECRET private URL')
        if case=='malformed':return '<broken'
        return xml(case)
    monkeypatch.setattr(probe,'fetch_xml',transport)
    return calls


def photo_request(c,f):
    documents=[upload(c,f),upload(c,f)]
    return {'official_mode':'MFDS_LIVE','normalization_confirmed':True,**{
        side:{'document_id':d['document_id'],'reviews':[{'index':n,'confirmed':True,'draft':draft(r['record'])} for n,r in enumerate(d['results'])]}
        for side,d in zip(('previous','current'),documents)}}


def hira(c,f,mode='MFDS_LIVE',code=None,empty=False):
    r=_row(1,'2026-10-01','가상알파정',code)
    r[6]='가상알파성분'
    import openpyxl
    w=openpyxl.load_workbook(BytesIO(_xlsx([] if empty else [r])))
    w.active.cell(2,13,'함량')
    if not empty:w.active.cell(3,13,'100mg')
    out=BytesIO();w.save(out);w.close();content=out.getvalue()
    return c.post(f'/api/dev-account/families/{f}/hira/process',data={'official_mode':mode},files={
        k:('virtual.xlsx',content,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet') for k in ('previous_file','current_file')})


@pytest.mark.parametrize('kind',['photo','hira'])
@pytest.mark.parametrize('case',['normal','partial','multiple','empty','unknown','mismatch','granule','name','ingredient','error','malformed'])
def test_offline_paths_normal_or_hold_no_mock_fallback(monkeypatch,kind,case,caplog):
    calls=fake(monkeypatch,case);c=account();f=family(c)
    response=c.post(f'/api/dev-account/families/{f}/photos/process',json=photo_request(c,f)) if kind=='photo' else hira(c,f)
    assert response.status_code==200,response.text
    result=response.json()['result']
    assert (result['mode'] if kind=='photo' else result['processing']['official_data_mode'])=='MFDS_LIVE'
    if kind=='hira':assert result['response_version']=='TASK_D_HIRA_MFDS_V1'
    diag=result['official_lookup']
    assert diag['mode']=='MFDS_LIVE' and diag['request_attempts']==1 and diag['cache_reuses']==1
    if case in ('error','malformed'):assert diag['queries'][0]['category'] in ('TIMEOUT_STAGE_UNDETERMINED','XML_INVALID')
    assert calls==[('가상알파정',1,3)]  # Both sides reuse the operation cache, no next page/retry.
    if case=='normal':
        assert result['reconciliation_result']['items'][0]['status']=='MAINTAINED'
        assert result['current_intake_states'][0]['verification_status']=='PENDING'
    else:
        assert result['reconciliation_result'] is None and result['current_intake_states']==[]
        assert result['stage']=='held' if kind=='photo' else result['comparison']['executed'] is False
    rows=result['normalization_results'] if kind=='photo' else result['accepted_results']
    for side in ('previous','current'):
        assert len(rows[side])==1
        if case in ('error','malformed','empty'):assert rows[side][0]['status']=='UNMATCHED'
        record=rows[side][0]['record'];assert record['product_name_raw']=='가상알파정'
        assert record['ingredients'][0]['product_strength']=='100'
        assert record['ingredients'][0]['product_strength_unit']=='mg'
        if case!='normal':assert record['mfds_product_id'] is None and rows[side][0]['status']!='MATCHED'
    assert 'SYNTHETIC-SECRET' not in response.text+caplog.text


def test_flag_off_no_request_even_with_key_no_fallback(monkeypatch):
    calls=fake(monkeypatch,'normal');c=account();f=family(c)
    monkeypatch.delenv('YAKEUM_MFDS_MATCH_ENABLED')
    body=photo_request(c,f)
    assert c.post(f'/api/dev-account/families/{f}/photos/process',json=body).status_code==503
    assert hira(c,f).status_code==503 and calls==[]
    # Legacy default is still the original Mock route.
    body.pop('official_mode');assert c.post(f'/api/dev-account/families/{f}/photos/process',json=body).json()['result']['mode']=='LOCAL_MOCK'
    assert hira(c,f,'LOCAL_MOCK').json()['result']['processing']['official_data_mode']=='LOCAL_MOCK'


@pytest.mark.parametrize('kind',['photo','hira'])
def test_persistence_sources_and_other_owner_access(monkeypatch,kind):
    calls=fake(monkeypatch,'normal');c=account();f=family(c);other=account('synthetic_b')
    data=(c.post(f'/api/dev-account/families/{f}/photos/process',json=photo_request(c,f)) if kind=='photo' else hira(c,f)).json()
    run=data['run_id'];base=f'/api/dev-account/families/{f}/'+('runs' if kind=='photo' else 'hira-runs')
    assert c.post(base+'/'+run+'/save',json={'intake_confirmations':[]}).status_code==200
    saved=c.get(base+'/'+run).json()['result']
    assert saved==data['result'] and len(calls)==1
    assert other.get(base+'/'+run).status_code==404
    assert other.delete(base+'/'+run).status_code==404
    if kind=='hira':assert c.delete(base+'/'+run).status_code==200


def test_budget_exhausted_for_another_product_no_second_call(monkeypatch):
    calls=fake(monkeypatch,'normal');lookup=LiveLookup('SYNTHETIC-SECRET')
    assert lookup.lookup('가상알파정')[0]=='COMPLETE'
    assert lookup.lookup('다른가상제품')[0]=='BUDGET_EXHAUSTED'
    assert len(calls)==1


def test_hira_insurance_unknown_not_inferred_and_zero_not_hold(monkeypatch):
    fake(monkeypatch,'normal');c=account();f=family(c)
    held=hira(c,f,code='000123456').json()['result']
    assert held['comparison']['executed'] is False and held['reconciliation_result'] is None
    assert held['accepted_results']['current'][0]['record']['drug_code']=='000123456'
    empty=hira(c,f,empty=True).json()['result']
    assert empty['comparison']=={'executed':True,'holds':[]} and empty['reconciliation_result']['items']==[]


def test_admin_once_ledger_not_modified_by_matching(monkeypatch):
    from backend.storage.local_db import database
    calls=fake(monkeypatch,'normal');c=account();f=family(c)
    with database() as db:
        db.execute(mfds_admin.SCHEMA)
        db.execute('INSERT INTO mfds_probe_attempts VALUES(?,?,?)',(mfds_admin.SLOT,'synthetic-time','{"old":true}'))
    assert hira(c,f).status_code==200
    with database() as db:assert db.execute('SELECT result FROM mfds_probe_attempts').fetchone()[0]=='{"old":true}'


def test_owner_session_csrf_and_origin_checked_before_live_lookup(monkeypatch):
    calls=fake(monkeypatch,'normal');c=account();f=family(c);body=photo_request(c,f);other=account('synthetic_other')
    root=f'/api/dev-account/families/{f}'
    assert other.post(root+'/photos/process',json=body).status_code==404
    assert hira(other,f).status_code==404
    assert c.post(root+'/photos/process',json=body,headers={'X-CSRF-Token':'wrong'}).status_code==403
    assert c.post(root+'/photos/process',json=body,headers={'Origin':'https://evil.invalid'}).status_code==403
    assert calls==[]


def test_source_missing_strength_and_wrong_basis_never_inferred(monkeypatch):
    calls=fake(monkeypatch,'normal');c=account();f=family(c);body=photo_request(c,f)
    for side in ('previous','current'):
        body[side]['reviews'][0]['draft']['ingredients'][0]['strength']=''
        body[side]['reviews'][0]['draft']['ingredients'][0]['unit']=''
    result=c.post(f'/api/dev-account/families/{f}/photos/process',json=body).json()['result']
    assert result['stage']=='held' and result['reconciliation_result'] is None
    assert result['normalization_results']['current'][0]['record']['ingredients'][0]['product_strength'] is None
    body=photo_request(c,f)
    for side in ('previous','current'):body[side]['reviews'][0]['draft']['doseUnit']='캡슐'
    result=c.post(f'/api/dev-account/families/{f}/photos/process',json=body).json()['result']
    assert result['stage']=='held'

"""Validate committed synthetic XLSX assets through actual uploads/parser/B/D; no DB or APIs."""
import hashlib,json
from pathlib import Path
from contextlib import ExitStack
from unittest.mock import patch
import openpyxl
import pytest
from fastapi.testclient import TestClient
from backend.api.main import app
from backend.services.hira.parser import parse_hira_excel

KIT=Path(__file__).resolve().parents[2]/'docs/testing/hira-badges'
MANIFEST=json.loads((KIT/'cases.json').read_text())
FAMILY='00000000-0000-4000-8000-000000000001'


def post_files(previous,current):
    with ExitStack() as stack:
        for target in ['backend.api.extraction_provider._request_gemini',
                       'backend.services.mfds.real_client.default_transport',
                       'backend.storage.local_db.database','backend.api.dev_account.database',
                       'backend.api.account_hira.database']:
            stack.enter_context(patch(target,side_effect=AssertionError('No live providers or database writes allowed')))
        return TestClient(app).post('/api/v2/reconciliations/hira',data={'family_member_id':FAMILY},files={
            side+'_file':(path.name,path.read_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            for side,path in [('previous',previous),('current',current)]})


def validate_case(case):
    for side,rel in case['files'].items():
        path=KIT/rel
        wb=openpyxl.load_workbook(path,read_only=True,data_only=False)
        try:
            assert '가상 시연 자료' in wb.active.cell(1,1).value
            assert wb.properties.creator=='YAKEUM SYNTHETIC DEMO'
            assert all(cell.data_type!='f' for row in wb.active for cell in row)
            if case['rows'][side]:assert wb.active.cell(3,8).value=='000123456'
        finally:wb.close()
        rows=parse_hira_excel(path)
        assert len(rows)==case['rows'][side]
    previous,current=(KIT/case['files'][s] for s in ('previous','current'))
    if not case['executed']:
        with patch('backend.api.hira_hold.reconcile_normalization_results',side_effect=AssertionError('Held inputs must not run comparison')):
            response=post_files(previous,current)
    else:response=post_files(previous,current)
    assert response.status_code==200,response.text
    body=response.json()
    assert body['processing']['file_processing']=='EXECUTED'
    assert body['processing']['official_data_mode']=='LOCAL_MOCK' and body['processing']['persisted'] is False
    assert body['comparison']['executed'] is case['executed']
    assert len(body['source_rows'])==sum(case['rows'].values())
    assert not body['review_results']['previous'] and not body['review_results']['current']
    for side in ('previous','current'):
        for candidate in body['accepted_results'][side]:
            record=candidate['record']
            assert record['drug_code']=='000123456' and record['user_confirmed'] is False
            assert record['ingredients'][0]['product_strength']=='500'
            assert record['ingredients'][0]['product_strength_unit']=='mg'
    if case['executed']:
        items=body['reconciliation_result']['items']
        assert [i['status'] for i in items]==case['statuses']
        assert [i['reason_code'] for i in items]==case['reason_codes']
        assert not body['comparison']['holds']
        if case['case']=='02_changed':
            assert items[0]['previous_record']['dose_per_administration']=='1'
            assert items[0]['current_record']['dose_per_administration']=='2'
        if case['case']=='03_new':assert items[0]['previous_record'] is None and items[0]['current_record'] is not None
        if case['case']=='04_previous_only':assert items[0]['previous_record'] is not None and items[0]['current_record'] is None
        assert all(s['intake_status']=='UNKNOWN' and s['verification_status']=='PENDING' for s in body['current_intake_states'])
    else:
        assert body['reconciliation_result'] is None and body['current_intake_states']==[]
        assert [h['kind'] for h in body['comparison']['holds']]==[case['hold_kind']]
        assert len(body['accepted_results']['previous'])==len(body['accepted_results']['current'])==1
        assert body['accepted_results']['current'][0]['record']['dose_per_administration'] is None
        assert body['source_rows'] and all(h['affected_refs'] for h in body['comparison']['holds'])
    return {'case':case['case'],'upload_http':response.status_code,'parser_rows':case['rows'],
            'executed':body['comparison']['executed'],'statuses':case['statuses'],
            'hold_kinds':[h['kind'] for h in body['comparison']['holds']],
            'source_rows':len(body['source_rows']),'originals_preserved':True,'external_api_calls':0,'database_calls':0}


@pytest.mark.parametrize('case',MANIFEST['cases'],ids=lambda c:c['case'])
def test_committed_demo_workbooks(case):validate_case(case)


def test_empty_comparison_is_not_a_hold_or_delete():
    response=post_files(KIT/'03_new/previous.xlsx',KIT/'04_previous_only/current.xlsx')
    assert response.status_code==200
    body=response.json()
    assert body['comparison']=={'executed':True,'holds':[]}
    assert body['reconciliation_result']['items']==[] and body['source_rows']==[]


def test_asset_checksums_match():
    for line in (KIT/'SHA256SUMS').read_text().splitlines():
        digest,relative=line.split('  ',1)
        assert hashlib.sha256((KIT/relative).read_bytes()).hexdigest()==digest
